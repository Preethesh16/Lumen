#!/usr/bin/env python3
"""Build the final Lumen deck by editing the supplied IBM template in place."""

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


TEMPLATE = Path("/home/infinity/Downloads/Lumen.pptx")
OUTPUT = Path("/home/infinity/Projects/Lumen/deliverables/Lumen-Project-Presentation.pptx")
SCREENSHOTS = Path("/home/infinity/Pictures/Screenshots/lumen ss")

DASHBOARD_URL = "https://lumen.64-227-166-157.sslip.io"
API_URL = "https://api.64-227-166-157.sslip.io/health"
GITHUB_URL = "https://github.com/Preethesh16/Lumen"


def text_shapes(slide):
    return [shape for shape in slide.shapes if shape.has_text_frame]


def picture_shapes(slide):
    return [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]


def set_paragraph(paragraph, text):
    for run in paragraph.runs:
        run.hyperlink.address = None
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run().text = text


def set_shape_paragraphs(shape, replacements):
    """Replace selected paragraph indexes while retaining template formatting."""
    paragraphs = shape.text_frame.paragraphs
    for index, text in replacements.items():
        set_paragraph(paragraphs[index], text)


def set_title(slide, text):
    title_shape = text_shapes(slide)[0]
    set_paragraph(title_shape.text_frame.paragraphs[0], text)


def fit_dimensions(image_path, left, top, width, height):
    with Image.open(image_path) as image:
        image_ratio = image.width / image.height
    box_ratio = width / height
    if image_ratio >= box_ratio:
        fitted_width = width
        fitted_height = int(width / image_ratio)
    else:
        fitted_height = height
        fitted_width = int(height * image_ratio)
    fitted_left = left + int((width - fitted_width) / 2)
    fitted_top = top + int((height - fitted_height) / 2)
    return fitted_left, fitted_top, fitted_width, fitted_height


def replace_picture(slide, picture, image_path):
    """Replace an illustration but keep its original box and z-order."""
    left, top, width, height = fit_dimensions(
        image_path, picture.left, picture.top, picture.width, picture.height
    )
    tree = picture._element.getparent()
    index = tree.index(picture._element)
    tree.remove(picture._element)
    replacement = slide.shapes.add_picture(str(image_path), left, top, width, height)
    tree.remove(replacement._element)
    tree.insert(index, replacement._element)
    return replacement


def add_contained_picture(slide, image_path, left, top, width, height):
    left, top, width, height = fit_dimensions(image_path, left, top, width, height)
    return slide.shapes.add_picture(str(image_path), left, top, width, height)


def set_hyperlinked_line(paragraph, text, address):
    set_paragraph(paragraph, text)
    paragraph.runs[0].hyperlink.address = address


def build():
    presentation = Presentation(str(TEMPLATE))
    slides = presentation.slides

    hero = SCREENSHOTS / "Screenshot From 2026-07-26 00-44-08.png"
    index = SCREENSHOTS / "Screenshot From 2026-07-26 00-44-19.png"
    detail = SCREENSHOTS / "Screenshot From 2026-07-26 00-44-42.png"
    evidence = SCREENSHOTS / "Screenshot From 2026-07-26 00-44-45.png"
    briefs = SCREENSHOTS / "Screenshot From 2026-07-26 00-44-49.png"
    telegram = SCREENSHOTS / "Screenshot From 2026-07-26 00-43-56.png"
    email = SCREENSHOTS / "Screenshot From 2026-07-26 02-38-38.png"

    # 1 — Cover and participant details.
    cover_shapes = text_shapes(slides[0])
    set_shape_paragraphs(
        cover_shapes[0],
        {0: "Lumen — Humanitarian Attention-Gap Intelligence"},
    )
    set_shape_paragraphs(
        cover_shapes[1],
        {
            0: "Full Name: Preethesh Carvalho, Deepthi C J",
            2: "Registered Mail ID: 23i61.preethesh@sjec.ac.in, 23i62.deepthi@sjec.ac.in",
            4: "College Name: St Joseph Engineering College, Mangalore",
            6: "Mobile: 8904145205, 9019932005",
            9: "IBM SkillsBuild Mail ID: 23i61.preethesh@sjec.ac.in, 23i62.deepthi@sjec.ac.in",
        },
    )

    # 2 — Introduction.
    set_title(slides[1], "Introduction:")
    set_shape_paragraphs(
        text_shapes(slides[1])[1],
        {
            0: "The Attention Gap:",
            1: "High-need crises can remain under-reported.",
            2: "Humanitarian and media data lives in separate systems.",
            3: "Lumen's Response:",
            4: "Compare need, coverage and underfunding transparently.",
            5: "From Evidence to Action:",
            6: "Rank 26 crises and deliver grounded briefs daily.",
        },
    )
    replace_picture(slides[1], picture_shapes(slides[1])[0], hero)

    # 3 — Problem and objectives.
    set_title(slides[2], "Problem Statement:")
    set_shape_paragraphs(
        text_shapes(slides[2])[1],
        {
            0: "Objective:",
            1: "Surface high-need, low-attention crises through transparent daily scoring.",
            2: "Specific Goals:",
            3: "Collect displacement, appeal, funding, disaster and coverage signals.",
            4: "Store raw observations before computing scores.",
            5: "Rank crises with a tested Attention-Gap formula.",
            6: "Explain trends and evidence in a public dashboard.",
            7: "Deliver briefs through Telegram and email.",
            8: "Remain useful when a source or AI provider fails.",
        },
    )
    replace_picture(slides[2], picture_shapes(slides[2])[0], index)

    # 4 — SDG alignment. Keep the template's gradient background picture.
    set_title(slides[3], "Sustainable Development Goals:")
    slide4_text = text_shapes(slides[3])
    set_shape_paragraphs(
        slide4_text[1],
        {
            0: "Goals Chosen: SDG 10, SDG 16 and SDG 17",
            1: "Problem Statement:",
            2: "Unequal crisis visibility weakens informed reporting, accountability and resource mobilisation.",
        },
    )
    set_shape_paragraphs(
        slide4_text[2],
        {
            0: "Rationale:",
            1: "Lumen supports fairer attention, explainable public-interest information, and partnerships across humanitarian data, journalism and funding actors.",
        },
    )
    replace_picture(slides[3], picture_shapes(slides[3])[1], detail)

    # 5 — Ingestion and preparation. Keep the template's gradient background picture.
    set_title(slides[4], "Data Collection and Preprocessing:")
    slide5_text = text_shapes(slides[4])
    set_shape_paragraphs(
        slide5_text[1],
        {
            0: "Data Sources:",
            1: "UNHCR displacement; OCHA FTS appeals and funding.",
            2: "Coverage and Disaster Signals:",
            3: "GDELT media volume; ReliefWeb active disasters.",
            4: "Daily adapters apply timeout, pacing, retry and validation.",
        },
    )
    set_shape_paragraphs(slide5_text[2], {0: "Evidence Storage and Preparation:"})
    set_shape_paragraphs(
        slide5_text[3],
        {
            0: "Observation Store:",
            1: "Raw readings are saved in PostgreSQL before scoring.",
            2: "Preparation:",
            3: "ISO3 mapping, parsing, log transform and missing-data handling.",
        },
    )
    replace_picture(slides[4], picture_shapes(slides[4])[1], evidence)

    # 6 — Scoring engine.
    set_title(slides[5], "Attention-Gap Scoring Engine:")
    set_shape_paragraphs(
        text_shapes(slides[5])[1],
        {
            0: "Algorithm:",
            1: "Need Score",
            2: "Displacement 50%, appeal size 30%, active disasters 20%.",
            3: "Missing components are dropped and remaining weights re-normalized.",
            4: "Coverage and Funding:",
            5: "GDELT estimates coverage; OCHA FTS supplies the funding gap.",
            6: "Formula v1.1.0:",
            7: "(need − coverage) × (1 + 0.5 × funding gap) + 0.3 × funding gap",
            8: "Versioned inputs make every result explainable and replayable.",
        },
    )
    replace_picture(slides[5], picture_shapes(slides[5])[0], detail)

    # 7 — Evaluation and reliability.
    set_title(slides[6], "Validation and Reliability:")
    set_shape_paragraphs(
        text_shapes(slides[6])[1],
        {
            0: "Verification:",
            1: "39 scoring unit tests cover normalization, missing data and edge cases.",
            2: "API integration tests run against PostgreSQL.",
            3: "CI checks tests, types, web build, migrations and committed secrets.",
            4: "Live-Data Result:",
            5: "UNHCR and OCHA FTS successfully ranked the 26-crisis cohort.",
            6: "Idempotent upserts prevent duplicate evidence during retries.",
            7: "Grounding Controls:",
            8: "Numeric facts are rendered directly from stored score inputs.",
            9: "Any AI error or invented number triggers a deterministic fallback.",
        },
    )
    replace_picture(slides[6], picture_shapes(slides[6])[0], index)

    # 8 — Stack and tools.
    set_title(slides[7], "Tools and Resources:")
    set_shape_paragraphs(
        text_shapes(slides[7])[1],
        {
            0: "Production Stack:",
            1: "Application:",
            2: "Next.js dashboard; Express and TypeScript API.",
            3: "Data and Orchestration:",
            4: "PostgreSQL 16; n8n daily scheduler; Docker Compose.",
            5: "Delivery and Operations:",
            6: "Optional Groq, Telegram, Resend, GitHub Actions, DigitalOcean and Caddy HTTPS.",
        },
    )
    replace_picture(slides[7], picture_shapes(slides[7])[0], hero)

    # 9 — Impact.
    set_title(slides[8], "Project Impact and Effectiveness:")
    set_shape_paragraphs(
        text_shapes(slides[8])[1],
        {
            0: "Problem Solving:",
            1: "For Journalists and Editors:",
            2: "Highlights overlooked stories and creates ready-to-share briefs.",
            3: "For NGOs and UN Teams:",
            4: "Reduces manual comparison and supports advocacy monitoring.",
            5: "For Donors and Policy Teams:",
            6: "Adds an explainable signal for attention and funding priorities.",
            7: "Expected Outcomes:",
            8: "Faster discovery, traceable communication and broader visibility.",
            9: "Measure freshness, delivery success, users, shares and actions.",
        },
    )
    replace_picture(slides[8], picture_shapes(slides[8])[0], briefs)

    # 10 — Why the approach is credible.
    set_title(slides[9], "Why It Will Work:")
    set_shape_paragraphs(
        text_shapes(slides[9])[1],
        {
            0: "Evidence-Driven:",
            1: "Combines trusted public sources and preserves every raw observation.",
            2: "Transparent by Design:",
            3: "Versioned scoring and score-linked briefs make outputs auditable.",
            4: "Resilient Operations:",
            5: "A failed source or AI provider cannot stop ranking and outreach.",
            6: "Proven End-to-End:",
            7: "Dashboard, API, database, n8n, Telegram and email are working live.",
        },
    )
    replace_picture(slides[9], picture_shapes(slides[9])[0], telegram)

    # 11 — Ranking result screenshot.
    set_title(slides[10], "Results and Output:")
    set_shape_paragraphs(
        text_shapes(slides[10])[1],
        {
            0: "The live ranking surfaces Syrian Arab Republic, Sudan, Yemen and Afghanistan as current high attention-gap crises.",
        },
    )
    add_contained_picture(
        slides[10], index, int(0.42 * 914400), int(0.82 * 914400),
        int(9.16 * 914400), int(3.58 * 914400)
    )

    # 12 — Public dashboard screenshot and link.
    set_title(slides[11], "Deployed Lumen Dashboard:")
    dashboard_caption = text_shapes(slides[11])[1]
    set_shape_paragraphs(
        dashboard_caption,
        {
            0: "Public HTTPS dashboard — explore rankings, crisis evidence, trends and grounded briefs.",
        },
    )
    dashboard_caption.text_frame.paragraphs[0].runs[0].hyperlink.address = DASHBOARD_URL
    deployed_picture = add_contained_picture(
        slides[11], hero, int(0.42 * 914400), int(0.82 * 914400),
        int(9.16 * 914400), int(3.58 * 914400)
    )
    deployed_picture.click_action.hyperlink.address = DASHBOARD_URL

    # 13 — Telegram and email delivery proof.
    set_title(slides[12], "Live Notification Delivery:")
    set_shape_paragraphs(
        text_shapes(slides[12])[1],
        {
            0: "Evidence-grounded humanitarian briefs were delivered successfully through Telegram and Resend email.",
        },
    )
    add_contained_picture(
        slides[12], telegram, int(0.35 * 914400), int(0.92 * 914400),
        int(4.35 * 914400), int(3.62 * 914400)
    )
    add_contained_picture(
        slides[12], email, int(4.88 * 914400), int(0.92 * 914400),
        int(4.77 * 914400), int(3.62 * 914400)
    )

    # 14 — Clickable project links. Keep the template's full-slide background.
    set_title(slides[13], "Project Links:")
    link_shape = text_shapes(slides[13])[1]
    link_paragraphs = link_shape.text_frame.paragraphs
    set_hyperlinked_line(link_paragraphs[0], f"GitHub Repository:  {GITHUB_URL}", GITHUB_URL)
    set_hyperlinked_line(link_paragraphs[1], f"Dashboard (Frontend):  {DASHBOARD_URL}", DASHBOARD_URL)
    set_hyperlinked_line(link_paragraphs[4], f"API Health (Backend):  {API_URL}", API_URL)
    for picture in picture_shapes(slides[13])[1:]:
        picture.click_action.hyperlink.address = (
            DASHBOARD_URL if picture.top < int(3 * 914400) else API_URL
        )

    # 15 — Conclusion.
    set_title(slides[14], "Conclusion:")
    set_shape_paragraphs(
        text_shapes(slides[14])[1],
        {
            0: "Summary:",
            1: "Lumen measures where humanitarian need outpaces attention.",
            2: "It connects daily public data, explainable scoring and live outreach.",
            3: "Future Work:",
            4: "Activate ReliefWeb after approval and calibrate with user feedback.",
            5: "Pilot with newsrooms and NGOs; add fixed-scale and sub-national signals.",
        },
    )
    replace_picture(slides[14], picture_shapes(slides[14])[0], briefs)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    build()
